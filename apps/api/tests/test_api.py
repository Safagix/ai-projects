from pathlib import Path
import time

import pytest
from fastapi.testclient import TestClient

from fashion_cad_api.config import settings


def test_create_laptop_bag_requests_physical_dimensions(client: TestClient) -> None:
    response = client.post("/api/projects", json={"name": "Bolso urbano", "product_type": "laptop_bag"})
    assert response.status_code == 201
    assert set(response.json()["missing"]) == {"laptop_width_mm", "laptop_height_mm", "laptop_depth_mm"}


def test_laptop_bag_exports_require_physical_dimensions(client: TestClient) -> None:
    created = client.post("/api/projects", json={"name": "Bolso incompleto", "product_type": "laptop_bag"}).json()
    for export_type in ("mockup", "pattern", "techpack"):
        response = client.post(f"/api/projects/{created['id']}/exports/{export_type}")
        assert response.status_code == 422
        assert "ancho" in response.json()["detail"].lower()


def test_local_brief_detects_bag_but_never_invents_laptop_dimensions(client: TestClient) -> None:
    response = client.post("/api/briefs/local", json={"description": "Bolso urbano reciclado para laptop de 16 pulgadas con bolsillo y cierre negro."})
    assert response.status_code == 200
    body = response.json()
    assert body["product_type"] == "laptop_bag"
    assert set(body["missing"]) == {"laptop_width_mm", "laptop_height_mm", "laptop_depth_mm"}
    assert "ancho" in body["next_question"].lower()


def test_mockup_export_is_a_binary_glb_with_revision_metadata(client: TestClient) -> None:
    created = client.post(
        "/api/projects",
        json={
            "name": "Bolso GLB",
            "product_type": "laptop_bag",
            "measurements_mm": {"laptop_width_mm": 355, "laptop_height_mm": 245, "laptop_depth_mm": 25},
            "components": ["bolsillo frontal", "cierre negro", "correa"],
        },
    ).json()
    response = client.post(f"/api/projects/{created['id']}/exports/mockup")
    assert response.status_code == 200
    glb_relative, metadata_relative = response.json()["artifacts"]
    assert (settings.artifacts_dir / glb_relative).read_bytes()[:4] == b"glTF"
    served_glb = client.get(f"/artifacts/{glb_relative}")
    assert served_glb.status_code == 200
    assert served_glb.content[:4] == b"glTF"
    metadata = (settings.artifacts_dir / metadata_relative).read_text(encoding="utf-8")
    assert '"revision": 1' in metadata
    assert '"parametric_local"' in metadata

def test_component_change_creates_new_pattern_revision_with_piece(client: TestClient) -> None:
    created = client.post(
        "/api/projects",
        json={
            "name": "Bolso sin bolsillo",
            "product_type": "laptop_bag",
            "measurements_mm": {"laptop_width_mm": 355, "laptop_height_mm": 245, "laptop_depth_mm": 25},
        },
    ).json()
    first = client.post(f"/api/projects/{created['id']}/exports/pattern").json()
    first_metadata = (settings.artifacts_dir / first["artifacts"][-1]).read_text(encoding="utf-8")
    assert "Front pocket" not in first_metadata

    changed = client.post(
        f"/api/projects/{created['id']}/operations",
        json={"kind": "add_component", "value": "bolsillo frontal"},
    )
    assert changed.status_code == 200
    second = client.post(f"/api/projects/{created['id']}/exports/pattern").json()
    second_metadata = (settings.artifacts_dir / second["artifacts"][-1]).read_text(encoding="utf-8")
    assert second["revision"] == 2
    assert "Front pocket" in second_metadata
    assert any(path.endswith("pattern-a0.pdf") for path in second["artifacts"])

def test_pattern_and_techpack_exports(client: TestClient) -> None:
    created = client.post(
        "/api/projects",
        json={
            "name": "Bolso 16 pulgadas",
            "product_type": "laptop_bag",
            "measurements_mm": {"laptop_width_mm": 355, "laptop_height_mm": 245, "laptop_depth_mm": 25},
            "materials": ["nylon reciclado"],
            "components": ["cierre negro"],
        },
    ).json()
    pattern = client.post(f"/api/projects/{created['id']}/exports/pattern")
    techpack = client.post(f"/api/projects/{created['id']}/exports/techpack")
    assert pattern.status_code == 200
    assert techpack.status_code == 200
    for relative_path in pattern.json()["artifacts"] + techpack.json()["artifacts"]:
        assert (settings.artifacts_dir / relative_path).exists()


def test_techpack_includes_factory_sheets(client: TestClient) -> None:
    from openpyxl import load_workbook

    created = client.post(
        "/api/projects",
        json={
            "name": "Camiseta tech pack",
            "product_type": "upper_garment",
            "measurements_mm": {"chest_mm": 1120, "body_length_mm": 720},
            "materials": ["algodón reciclado"],
            "components": ["cuello rib"],
        },
    ).json()
    response = client.post(f"/api/projects/{created['id']}/exports/techpack")
    assert response.status_code == 200
    workbook_path = settings.artifacts_dir / response.json()["artifacts"][1]
    workbook = load_workbook(workbook_path, read_only=True)
    assert {"Summary", "BOM", "POM", "Construction", "Grading", "Labels_Packaging", "Material_Evidence", "Revision_History"} <= set(workbook.sheetnames)

def test_semantic_rag_never_falls_back_to_fake_vectors_without_bge_model(client: TestClient) -> None:
    client.post("/api/rag/documents", json={"title": "Nylon", "content": "Nylon reciclado.", "source": "manual"})
    response = client.post("/api/rag/semantic/reindex")
    assert response.status_code == 503
    assert "BGE-M3" in response.json()["detail"]


def test_local_assistant_applies_chat_instructions_as_revisions(client: TestClient) -> None:
    created = client.post("/api/projects", json={"name": "Bolso chat", "product_type": "laptop_bag"}).json()
    response = client.post(
        f"/api/projects/{created['id']}/assistant",
        json={"message": "Agregá bolsillo, cierre y correa; usá nylon reciclado para 355 × 245 × 25 mm."},
    )
    assert response.status_code == 200
    body = response.json()
    assert {"bolsillo frontal", "cierre negro", "correa"} <= set(body["design"]["components"])
    assert "nylon reciclado" in body["design"]["materials"]
    assert body["design"]["measurements_mm"] == {
        "chest_mm": None, "body_length_mm": None,
        "laptop_width_mm": 355.0, "laptop_height_mm": 245.0, "laptop_depth_mm": 25.0,
    }
    assert body["design"]["revision"] == 8


def test_operation_persists_selected_mode(client: TestClient) -> None:
    created = client.post("/api/projects", json={"name": "Modo", "product_type": "upper_garment"}).json()
    response = client.post(f"/api/projects/{created['id']}/operations", json={"kind": "set_mode", "value": "hybrid"})
    assert response.status_code == 200
    assert response.json()["mode"] == "hybrid"


def test_semantic_reindex_runs_as_a_job_and_rebuilds_without_duplicates(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    from fashion_cad_api import main, vector_store

    class FakeEmbedder:
        batches: list[int] = []

        def ensure_available(self) -> None:
            return None

        def embed_documents(self, texts: list[str]) -> list[list[float]]:
            self.batches.append(len(texts))
            return [[float(index + 1), 0.5, 0.25] for index, _ in enumerate(texts)]

        def embed_query(self, _text: str) -> list[float]:
            return [1.0, 0.5, 0.25]

    monkeypatch.setattr(vector_store, "BgeM3Embedder", FakeEmbedder)
    for number in range(5):
        client.post(
            "/api/rag/documents",
            json={"title": f"Nylon {number}", "content": f"Nylon reciclado para bolso {number}.", "source": "manual"},
        )

    first = client.post("/api/rag/semantic/reindex")
    assert first.status_code == 202
    first_job = first.json()
    assert first_job["kind"] == "embedding_index"

    for _ in range(40):
        status = client.get(f"/api/jobs/{first_job['id']}").json()
        if status["status"] in {"completed", "failed", "cancelled"}:
            break
        time.sleep(0.025)
    assert status["status"] == "completed", status
    assert status["progress"] == status["total"] == 5
    assert FakeEmbedder.batches == [4, 1]

    result = client.get("/api/rag/semantic/search", params={"query": "nylon bolso"})
    assert result.status_code == 200
    assert result.json()[0]["source"] == "manual"
    assert result.json()[0]["chunk_id"]
    assert result.json()[0]["chunk_number"] == 1

    second = client.post("/api/rag/semantic/reindex")
    assert second.status_code == 202
    for _ in range(40):
        status = client.get(f"/api/jobs/{second.json()['id']}").json()
        if status["status"] in {"completed", "failed", "cancelled"}:
            break
        time.sleep(0.025)
    assert status["status"] == "completed"
    state = main.repository.get_semantic_index()
    assert state and state["source_count"] == 5
    assert vector_store.LanceSemanticStore(table_name=state["table_name"]).count() == 5


def test_queued_embedding_job_can_be_cancelled(client: TestClient) -> None:
    created = client.post("/api/jobs", json={"kind": "embedding_index"})
    assert created.status_code == 201
    cancelled = client.delete(f"/api/jobs/{created.json()['id']}")
    assert cancelled.status_code == 200
    assert cancelled.json()["status"] == "cancelled"

def test_rag_fts_returns_source(client: TestClient) -> None:
    client.post("/api/rag/documents", json={"title": "Nylon reciclado", "content": "El nylon reciclado requiere evidencia de origen.", "source": "material-sheet.pdf"})
    response = client.get("/api/rag/search", params={"query": "nylon reciclado"})
    assert response.status_code == 200
    assert response.json()[0]["source"] == "material-sheet.pdf"


def test_rag_fts_returns_ocr_confidence_when_imported_with_ocr(client: TestClient) -> None:
    from fashion_cad_api import main

    main.repository.add_rag_document(
        "Página OCR",
        "Puntada de seguridad para bolso.",
        "library/scan.pdf",
        source_page=4,
        chunk_number=1,
        ocr_confidence=0.9234,
    )
    response = client.get("/api/rag/search", params={"query": "puntada seguridad"})
    assert response.status_code == 200
    assert response.json()[0]["ocr_confidence"] == 0.9234


def test_rag_local_import_is_confined_to_library(client: TestClient) -> None:
    library = settings.library_dir
    library.mkdir(parents=True)
    (library / "nylon.md").write_text("Nylon reciclado con certificado de origen.", encoding="utf-8")

    imported = client.post("/api/rag/import", json={"relative_path": "nylon.md"})
    assert imported.status_code == 201
    assert imported.json()["source"] == "library/nylon.md"

    search = client.get("/api/rag/search", params={"query": "certificado origen"})
    assert search.status_code == 200
    assert search.json()[0]["source"] == "library/nylon.md"

    escaped = client.post("/api/rag/import", json={"relative_path": "..\\secret.txt"})
    assert escaped.status_code == 422


def test_rag_reimport_replaces_prior_chunks_from_the_same_source(client: TestClient) -> None:
    from fashion_cad_api import main

    library = settings.library_dir
    library.mkdir(parents=True, exist_ok=True)
    source = library / "reimport.md"
    source.write_text("Nylon reciclado para el primer bolso.", encoding="utf-8")
    assert client.post("/api/rag/import", json={"relative_path": "reimport.md"}).status_code == 201
    source.write_text("RPET 600D para el segundo bolso.", encoding="utf-8")
    assert client.post("/api/rag/import", json={"relative_path": "reimport.md"}).status_code == 201
    rows = [row for row in main.repository.list_rag_documents() if row["source"] == "library/reimport.md"]
    assert len(rows) == 1
    assert "RPET" in rows[0]["content"]

def test_rag_local_pdf_import_extracts_text(client: TestClient) -> None:
    from reportlab.pdfgen.canvas import Canvas

    library = settings.library_dir
    library.mkdir(parents=True, exist_ok=True)
    pdf_path = library / "construction.pdf"
    canvas = Canvas(str(pdf_path))
    canvas.drawString(72, 720, "La costura lateral usa cinco puntadas por centimetro.")
    canvas.save()

    imported = client.post("/api/rag/import", json={"relative_path": "construction.pdf"})
    assert imported.status_code == 201
    search = client.get("/api/rag/search", params={"query": "costura lateral"})
    assert search.status_code == 200
    assert search.json()[0]["source"] == "library/construction.pdf"

def test_cloud_consent_is_audited_and_local_mode_is_rejected(client: TestClient) -> None:
    local = client.post("/api/projects", json={"name": "Local", "product_type": "upper_garment"}).json()
    denied = client.post(
        f"/api/projects/{local['id']}/cloud-consents",
        json={"asset_id": "design_document", "provider": "openai", "purpose": "interpret sketch", "estimated_cost_usd": 0.03, "approved": True},
    )
    assert denied.status_code == 422

    hybrid = client.post("/api/projects", json={"name": "Hybrid", "product_type": "upper_garment", "mode": "hybrid"}).json()
    approved = client.post(
        f"/api/projects/{hybrid['id']}/cloud-consents",
        json={"asset_id": "design_document", "provider": "openai", "purpose": "interpret sketch", "estimated_cost_usd": 0.03, "approved": True},
    )
    assert approved.status_code == 201
    audit = client.get(f"/api/projects/{hybrid['id']}/cloud-consents")
    assert audit.status_code == 200
    assert audit.json()[0]["approved"] is True

def test_only_one_gpu_exclusive_job_can_be_queued(client: TestClient) -> None:
    first = client.post("/api/jobs", json={"kind": "blender_render", "gpu_exclusive": True})
    assert first.status_code == 201
    status = client.get(f"/api/jobs/{first.json()['id']}")
    assert status.status_code == 200
    assert status.json()["status"] == "queued"

    second = client.post("/api/jobs", json={"kind": "model_inference", "gpu_exclusive": True})
    assert second.status_code == 409

    cpu_job = client.post("/api/jobs", json={"kind": "embedding_index", "gpu_exclusive": False})
    assert cpu_job.status_code == 201

def test_operator_is_disabled_by_default(client: TestClient) -> None:
    response = client.post("/api/operator/sessions", json={"requested_windows": ["Fashion CAD Studio"]})
    assert response.status_code == 403


def test_operator_persists_audit_pauses_on_focus_loss_and_requires_confirmation(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The safety boundary is testable without driving a real desktop."""
    from fashion_cad_api import main

    class FakePyAutoGui:
        calls: list[tuple[str, object]] = []

        @classmethod
        def click(cls, x: int, y: int) -> None:
            cls.calls.append(("click", (x, y)))

        @classmethod
        def press(cls, key: str) -> None:
            cls.calls.append(("press", key))

        @classmethod
        def write(cls, text: str, interval: float) -> None:
            cls.calls.append(("write", text))

    from fashion_cad_api.config import settings

    settings.operator_enabled = True
    main.operator_guard.stop_all(reason="test cleanup")
    monkeypatch.setattr(main.operator_guard, "foreground_window", lambda: "Fashion CAD Studio")
    monkeypatch.setattr(main.operator_guard, "window_bounds", lambda: (0, 0, 1600, 900))
    monkeypatch.setattr(main.operator_guard, "automation", lambda: FakePyAutoGui)
    monkeypatch.setattr(main.operator_guard._overlay, "show", lambda _text: None)
    monkeypatch.setattr(main.operator_guard._overlay, "close", lambda: None)
    monkeypatch.setattr(main.operator_guard._kill_switch, "start", lambda _callback: True)

    started = client.post("/api/operator/sessions", json={"requested_windows": ["Fashion CAD Studio"]})
    assert started.status_code == 200
    session_id = started.json()["session_id"]

    clicked = client.post(
        "/api/operator/click",
        json={"session_id": session_id, "window_name": "Fashion CAD Studio", "x": 12, "y": 18},
    )
    assert clicked.status_code == 200
    assert FakePyAutoGui.calls == [("click", (12, 18))]

    blocked = client.post(
        "/api/operator/type",
        json={
            "session_id": session_id,
            "window_name": "Fashion CAD Studio",
            "text": "export",
            "intent": "export",
        },
    )
    assert blocked.status_code == 428

    confirmation = client.post(
        "/api/operator/confirmations",
        json={
            "session_id": session_id,
            "window_name": "Fashion CAD Studio",
            "intent": "export",
            "summary": "Exportar mockup GLB de revisión 1",
        },
    )
    assert confirmation.status_code == 201
    approved = client.post(
        "/api/operator/type",
        json={
            "session_id": session_id,
            "window_name": "Fashion CAD Studio",
            "text": "export",
            "intent": "export",
            "confirmation_id": confirmation.json()["confirmation_id"],
        },
    )
    assert approved.status_code == 200

    monkeypatch.setattr(main.operator_guard, "foreground_window", lambda: "")
    paused = client.post(
        "/api/operator/keypress",
        json={"session_id": session_id, "window_name": "Fashion CAD Studio", "key": "enter"},
    )
    assert paused.status_code == 409
    assert "pausada" in paused.json()["detail"].lower()

    audit = client.get(f"/api/operator/sessions/{session_id}/audit")
    assert audit.status_code == 200
    actions = [entry["action"] for entry in audit.json()]
    assert {"session_started", "click", "confirmation_requested", "type", "focus_lost"} <= set(actions)
