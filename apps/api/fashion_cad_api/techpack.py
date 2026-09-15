from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas

from .config import settings

DRAFT = "Por confirmar en muestra"


def _tolerance(document: dict[str, Any]) -> str:
    return "± 5 mm" if document["product_type"] == "laptop_bag" else "± 10 mm"


def _component_type(value: str) -> str:
    normalized = value.lower()
    if "cierre" in normalized or "zipper" in normalized:
        return "Hardware / zipper"
    if "correa" in normalized or "asa" in normalized or "strap" in normalized:
        return "Hardware / webbing"
    if "bolsillo" in normalized or "pocket" in normalized:
        return "Component / pocket"
    if "cuello" in normalized or "collar" in normalized:
        return "Component / collar"
    return "Component"


def _bom_rows(document: dict[str, Any]) -> list[list[str]]:
    rows: list[list[str]] = []
    for material in document["materials"]:
        evidence = "Evidencia RAG requerida" if "recicl" in material.lower() else DRAFT
        rows.append(["Material", material, DRAFT, evidence])
    for component in document["components"]:
        rows.append([_component_type(component), component, DRAFT, DRAFT])
    return rows or [["Material", "Sin material definido", DRAFT, DRAFT]]


def _pom_rows(document: dict[str, Any]) -> list[list[str]]:
    return [[name, str(value), "mm", _tolerance(document)] for name, value in document["measurements_mm"].items() if value]


def _construction_rows(document: dict[str, Any]) -> list[list[str]]:
    base = "Costura lateral y remate" if document["product_type"] == "upper_garment" else "Costura perimetral, forro y protección de bordes"
    rows = [["Base", base, "SPI y tipo de puntada: " + DRAFT]]
    for component in document["components"]:
        rows.append([component, "Ubicación y refuerzo: " + DRAFT, "Validar en muestra"])
    return rows


def export_techpack(document: dict[str, Any]) -> list[str]:
    base = settings.artifacts_dir / "techpacks" / document["id"] / f"rev-{document['revision']}"
    base.mkdir(parents=True, exist_ok=True)
    pdf_path, workbook_path = base / "tech-pack.pdf", base / "tech-pack.xlsx"
    _write_pdf(document, pdf_path)
    _write_workbook(document, workbook_path)
    return [str(pdf_path.relative_to(settings.artifacts_dir)), str(workbook_path.relative_to(settings.artifacts_dir))]


def _new_page(pdf: canvas.Canvas, title: str, document: dict[str, Any]) -> float:
    width, height = A4
    pdf.setFillColorRGB(0.04, 0.05, 0.07)
    pdf.rect(0, 0, width, height, fill=1, stroke=0)
    pdf.setFillColorRGB(0.95, 0.91, 0.82)
    pdf.setFont("Helvetica-Bold", 17)
    pdf.drawString(18 * mm, 276 * mm, "FASHION CAD STUDIO")
    pdf.setFont("Helvetica", 8)
    pdf.drawRightString(192 * mm, 276 * mm, f"{document['name']} · REV {document['revision']}")
    pdf.setFont("Helvetica-Bold", 13)
    pdf.drawString(18 * mm, 260 * mm, title)
    return 246


def _write_lines(pdf: canvas.Canvas, rows: list[list[str]], headings: list[str], y: float) -> None:
    pdf.setFont("Helvetica-Bold", 8)
    pdf.setFillColorRGB(0.95, 0.91, 0.82)
    pdf.drawString(18 * mm, y * mm, "  |  ".join(headings))
    y -= 8
    pdf.setFont("Helvetica", 8)
    pdf.setFillColorRGB(0.9, 0.9, 0.9)
    for row in rows:
        line = "  |  ".join(str(cell) for cell in row)
        pdf.drawString(18 * mm, y * mm, line[:140])
        y -= 7
        if y < 25:
            pdf.showPage()
            y = 265


def _write_pdf(document: dict[str, Any], path: Path) -> None:
    pdf = canvas.Canvas(str(path), pagesize=A4)
    y = _new_page(pdf, "Resumen técnico", document)
    pdf.setFillColorRGB(0.95, 0.91, 0.82)
    pdf.setFont("Helvetica-Bold", 20)
    pdf.drawString(18 * mm, y * mm, document["name"])
    y -= 16
    pdf.setFont("Helvetica", 9)
    lines = [
        f"SKU: {DRAFT}",
        f"Producto: {document['product_type']}",
        f"Estado: {document['approval_state']}",
        f"Confianza: {document['confidence']}",
        "Flat técnico: mockup paramétrico y patrón asociados a esta revisión.",
        "Todas las medidas, tolerancias y materiales requieren aprobación física antes de producción.",
    ]
    for line in lines:
        pdf.drawString(18 * mm, y * mm, line)
        y -= 8
    pdf.showPage()

    y = _new_page(pdf, "BOM · Bill of Materials", document)
    _write_lines(pdf, _bom_rows(document), ["Tipo", "Material / avío", "Especificación", "Evidencia"], y)
    pdf.showPage()

    y = _new_page(pdf, "POM · Points of Measure", document)
    _write_lines(pdf, _pom_rows(document) or [["Sin medidas", "", "", DRAFT]], ["POM", "Valor", "Unidad", "Tolerancia"], y)
    pdf.showPage()

    y = _new_page(pdf, "Construcción, grading y empaque", document)
    _write_lines(pdf, _construction_rows(document), ["Operación", "Construcción", "Control"], y)
    pdf.setFillColorRGB(0.95, 0.91, 0.82)
    pdf.setFont("Helvetica-Bold", 9)
    pdf.drawString(18 * mm, 100 * mm, "GRADING")
    pdf.setFont("Helvetica", 8)
    pdf.drawString(18 * mm, 92 * mm, "Base: " + ("Talla M" if document["product_type"] == "upper_garment" else "Talla única") + " · Incrementos: " + DRAFT)
    pdf.drawString(18 * mm, 80 * mm, "ETIQUETAS Y EMPAQUE")
    pdf.drawString(18 * mm, 72 * mm, "Composición, origen, cuidado, bolsa y etiqueta colgante: " + DRAFT)
    pdf.setFillColorRGB(0.72, 0.74, 0.78)
    pdf.drawString(18 * mm, 20 * mm, "Borrador técnico. Validar muestra, tolerancias, SPI, avíos y certificaciones antes de producción.")
    pdf.save()


def _style_sheet(sheet) -> None:
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="20312D")
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 14), 48)


def _write_workbook(document: dict[str, Any], path: Path) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Field", "Value"])
    for key in ["name", "product_type", "revision", "approval_state", "confidence", "mode"]:
        summary.append([key, document[key]])
    summary.append(["SKU", DRAFT])

    bom = workbook.create_sheet("BOM")
    bom.append(["Type", "Material / component", "Specification", "Evidence"])
    for row in _bom_rows(document):
        bom.append(row)

    pom = workbook.create_sheet("POM")
    pom.append(["Measurement", "Value", "Unit", "Tolerance"])
    for row in _pom_rows(document):
        pom.append(row)

    construction = workbook.create_sheet("Construction")
    construction.append(["Operation", "Construction", "Control"])
    for row in _construction_rows(document):
        construction.append(row)

    grading = workbook.create_sheet("Grading")
    grading.append(["Base size", "POM", "Increment", "Status"])
    grading.append(["M" if document["product_type"] == "upper_garment" else "One size", "All", DRAFT, "Not approved"])

    labels = workbook.create_sheet("Labels_Packaging")
    labels.append(["Item", "Specification", "Status"])
    for item in ["Care label", "Composition label", "Hangtag", "Polybag/carton"]:
        labels.append([item, DRAFT, "Required before production"])

    evidence = workbook.create_sheet("Material_Evidence")
    evidence.append(["Material", "Claim", "Source / certification", "Status"])
    for material in document["materials"] or ["No material defined"]:
        claim = "Recycled claim" if "recicl" in material.lower() else "No claim"
        evidence.append([material, claim, "RAG citation required", "Pending"])

    revisions = workbook.create_sheet("Revision_History")
    revisions.append(["Revision", "Timestamp", "State", "Comment"])
    revisions.append([document["revision"], document["updated_at"], document["approval_state"], "Current DesignDocument revision"])

    for sheet in workbook.worksheets:
        _style_sheet(sheet)
    workbook.save(path)